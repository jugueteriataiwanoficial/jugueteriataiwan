#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
sync_posgold.py — Carga automatica de los 300 productos mas rentables de
POSGold al index.html de Jugueteria Taiwan (jugueteriataiwan.com).

Como funciona (cron: programar_tarea.ps1 / Task Scheduler de Windows):
  1. Carga credenciales de .env — JAMAS van al repo de GitHub
  2. Genera JWT fresco: POST /apiGold/Login_Api/LoginJWTToken
  3. Descarga productos: GET /apiGold/ProductoApi/GetProduct_V6
     (bodega configurada en .env, por defecto la 2)
  4. Carga costos del ultimo Excel exportado de POSGold (por Productoid)
  5. FILTROS DUROS (regla de la tienda): activo + Disponible>0 en bodega
     + foto real en POSGold + categoria vendible + costo conocido
  6. Rentabilidad sobre el precio detal:
       precio_base_sin_iva = Precio1 / (1 + IVA/100)
       margen_costo  = (precio_base_sin_iva - costo) / costo     (default)
       margen_precio = (precio_base_sin_iva - costo) / precio_base_sin_iva
       utilidad      = precio_base_sin_iva - costo
  7. Toma el TOP N (los destacados que cumplen entran primero si
     INCLUIR_DESTACADOS=true) y resuelve foto + pagina de cada uno
     (cache data/mapa_imagenes.json + busqueda en la web, verificando
     que el SKU de la pagina sea el codigo de POSGold)
  8. Reemplaza el bloque /*__TOP300:INICIO__*/ ... /*__TOP300:FIN__*/
     del index.html — el producto que deja de cumplir se oculta solo
  9. Reporte local en reportes/ (contiene costos: NO publicar)
 10. Con --publicar: git commit + push -> GitHub Pages republica solo

Uso:
  python sync_posgold.py --dry-run     (prueba sin tocar el index)
  python sync_posgold.py               (carga el index, sin publicar)
  python sync_posgold.py --publicar    (carga y publica en GitHub)
"""

import argparse
import csv
import datetime as dt
import glob
import io
import json
import os
import re
import shutil
import ssl
import subprocess
import sys
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request

# ------------------------------------------------------------------
# Configuracion general
# ------------------------------------------------------------------

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
INDEX = os.path.join(BASE_DIR, "index.html")
DIR_DATA = os.path.join(BASE_DIR, "data")
DIR_REPORTES = os.path.join(BASE_DIR, "reportes")
CACHE_IMAGENES = os.path.join(DIR_DATA, "mapa_imagenes.json")
ENV_FILE = os.path.join(BASE_DIR, ".env")

MARCA_INICIO = "/*__TOP300:INICIO__*/"
MARCA_FIN = "/*__TOP300:FIN__*/"
PLACEHOLDER_POSGOLD = "no-disponible"  # nombre de foto "sin imagen" de POSGold

# El certificado del servidor POSGold no siempre cuadra en Windows y el
# script solo hace consultas de lectura: se desactiva la verificacion.
CTX = ssl.create_default_context()
CTX.check_hostname = False
CTX.verify_mode = ssl.CERT_NONE
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) sync-posgold/1.0"

# Categorias de POSGold -> slugs de categoria del index.html
CATEGORIAS_POSGOLD = {
    "JUGUETERIA": "jugueteria",
    "COLECCION": "colecciones",
    "CACHARRO": "cacharro",
    "SLIME": "slime",
    "PINATERIA": "piñateria",
    "JUEGO DE MESA": "juego-de-mesa",
    "BLOQUES DE CONSTRUCCION": "armatodo",
    "DIDACTICO": "didactico",
    "CARTONERIA": "piñateria",
    "DEPORTIVO": "deporte",
    "BOLSO Y PELUCHES": "peluches",
    "BEBES": "bebe",
    "ESCOLAR": "escolar",
    "TERMOS": "termos",
    "CUBO RUBIK": "cubo-rubik",
    "BURBUJEROS": "burbujeros",
    "ANIMALES DE CAUCHO EN BOLSA": "animales",
    "ARMATODO": "armatodo",
    "LLAVEROS": "jugueteria",
    "MADEROTERAPIA": "hogar",
    "DINOSAURIO": "animales",
}
# Categorias que no son productos de venta (se ocultan siempre)
CATEGORIAS_NO_VENDIBLES = {"HONORARIOS Y SERVICIOS", "ARRENDAMIENTOS"}

# Fallback por palabras clave cuando la categoria exacta no esta mapeada
CLAVES_CATEGORIA = [
    ("PELUCH", "peluches"),
    ("CARRO", "vehiculos"),
    ("VEHICUL", "vehiculos"),
    ("RELOJ", "relojeria"),
    ("TERMO", "termos"),
    ("RUBIK", "cubo-rubik"),
    ("CUBO", "cubo-rubik"),
    ("ANIMAL", "animales"),
    ("ESCOLAR", "escolar"),
    ("DIDACT", "didactico"),
    ("JUEGO DE MESA", "juego-de-mesa"),
    ("BURBUJA", "burbujeros"),
    ("BURBUJER", "burbujeros"),
    ("SLIME", "slime"),
    ("ARENA", "slime"),
    ("HIDROGEL", "slime"),
    ("DEPORT", "deporte"),
    ("PINATA", "piñateria"),
    ("CACHARR", "cacharro"),
    ("COLECC", "colecciones"),
    ("FIGURA", "colecciones"),
    ("MUNECA", "colecciones"),
    ("ANIME", "colecciones"),
    ("BLOQUE", "armatodo"),
    ("ARMATODO", "armatodo"),
    ("ARMABLE", "armatodo"),
    ("BEBE", "bebe"),
    ("TECNOL", "tecnologia"),
    ("AUDIFON", "tecnologia"),
    ("TECLADO", "tecnologia"),
    ("HOGAR", "hogar"),
    ("COCINA", "hogar"),
]

# ------------------------------------------------------------------
# Utilidades
# ------------------------------------------------------------------


def log(msg):
    print(msg, flush=True)


def sin_acentos(txt):
    txt = unicodedata.normalize("NFD", txt or "")
    return "".join(c for c in txt if unicodedata.category(c) != "Mn").upper()


def cargar_env(ruta=ENV_FILE):
    env = {}
    if not os.path.exists(ruta):
        log("ERROR: falta el archivo .env (credenciales POSGold).")
        log("       Crea .env en la raiz del proyecto con POSGOLD_BASE,")
        log("       POSGOLD_TOKEN_USER y POSGOLD_TOKEN_PASS (ver info taiwan.txt).")
        sys.exit(1)
    with open(ruta, encoding="utf-8-sig") as f:
        for linea in f:
            linea = linea.strip()
            if not linea or linea.startswith("#") or "=" not in linea:
                continue
            k, v = linea.split("=", 1)
            env[k.strip()] = v.strip().strip('"').strip("'")
    return env


def http_texto(url, headers=None, timeout=90, metodo="GET", datos=None, reintentos=3):
    """GET/POST que devuelve texto. Reintenta errores de red y 5xx."""
    hs = {"User-Agent": UA}
    if headers:
        hs.update(headers)
    ultimo_error = None
    for intento in range(reintentos):
        req = urllib.request.Request(url, headers=hs, method=metodo,
                                     data=datos.encode("utf-8") if datos else None)
        try:
            with urllib.request.urlopen(req, timeout=timeout, context=CTX) as r:
                return r.status, r.read().decode("utf-8", errors="replace")
        except urllib.error.HTTPError as e:
            cuerpo = ""
            try:
                cuerpo = e.read().decode("utf-8", errors="replace")[:300]
            except Exception:
                pass
            if e.code >= 500 and intento < reintentos - 1:
                time.sleep(2 * (intento + 1))
                continue
            return e.code, cuerpo
        except Exception as e:  # timeout, conexion caida...
            ultimo_error = e
            if intento < reintentos - 1:
                time.sleep(2 * (intento + 1))
                continue
            raise
    raise RuntimeError("http_texto fallo: " + str(ultimo_error))


def imagen_responde(url):
    """HEAD con respaldo GET-Range: confirma que la foto existe en la web.
    Devuelve (ok, codigo_http). Codigo 0 = error de red o timeout."""
    req = urllib.request.Request(url, method="HEAD", headers={"User-Agent": UA})
    try:
        with urllib.request.urlopen(req, timeout=25, context=CTX) as r:
            return r.status == 200, r.status
    except urllib.error.HTTPError as e:
        if e.code in (403, 405, 501):  # servidor sin HEAD -> GET parcial
            req = urllib.request.Request(url, method="GET",
                                        headers={"User-Agent": UA, "Range": "bytes=0-1"})
            try:
                with urllib.request.urlopen(req, timeout=25, context=CTX) as r:
                    return r.status in (200, 206), r.status
            except urllib.error.HTTPError as e2:
                return False, e2.code
            except Exception:
                return False, 0
        return False, e.code
    except Exception:
        return False, 0


# ------------------------------------------------------------------
# PASO 1-2: JWT fresco
# ------------------------------------------------------------------


def login_jwt(env):
    qs = urllib.parse.urlencode({"user": env["POSGOLD_TOKEN_USER"],
                                "Password": env["POSGOLD_TOKEN_PASS"]})
    url = f"{env['POSGOLD_BASE']}/apiGold/Login_Api/LoginJWTToken?{qs}"
    codigo, texto = http_texto(url, metodo="POST", datos="{}", timeout=60)
    if codigo != 200:
        log(f"ERROR: login JWT respondio {codigo}: {texto[:200]}")
        sys.exit(1)
    data = json.loads(texto)
    if not data.get("Status") or not data.get("Token"):
        log("ERROR: POSGold no entrego token: " + str(data)[:200])
        sys.exit(1)
    return data["Token"]


# ------------------------------------------------------------------
# PASO 3: productos de POSGold (bodega configurada)
# ------------------------------------------------------------------


def consultar_productos(env, jwt):
    productos = []
    pagina = 1
    total = None
    while True:
        params = {
            "empresaid": env.get("POSGOLD_EMPRESAID", "1"),
            "usuarioid": env.get("POSGOLD_USUARIOID", "1"),
            "codigo": "", "descripcion": "", "categoriaid": "", "grupoid": "",
            "inicial": "0", "items_x_pagina": "50000", "pagina": str(pagina),
            "subgrupoid": "", "precio_min": "", "precio_max": "",
            "bodegaid": env.get("POSGOLD_BODEGAID", "2"),
            "orden": "", "api": "1", "api2": "1", "nombreFull": "",
            "ult_mov": "2000-01-01", "activo": "",
        }
        url = f"{env['POSGOLD_BASE']}/apiGold/ProductoApi/GetProduct_V6?{urllib.parse.urlencode(params)}"
        codigo, texto = http_texto(url, headers={"Authorization": "Bearer " + jwt}, timeout=180)
        if codigo != 200:
            log(f"ERROR: consulta de productos respondio {codigo}: {texto[:200]}")
            sys.exit(1)
        data = json.loads(texto)
        if not data.get("Status"):
            log("ERROR: POSGold rechazo la consulta: " + str(data.get("Msj"))[:200])
            sys.exit(1)
        lote = data.get("Datos") or []
        productos.extend(lote)
        total = data.get("Conteo", len(productos))
        log(f"  pagina {pagina}: {len(lote)} productos (total declarado: {total})")
        if len(productos) >= total or not lote:
            break
        pagina += 1
    return productos


# ------------------------------------------------------------------
# PASO 4: costos del ultimo Excel exportado de POSGold
# ------------------------------------------------------------------


def encontrar_excel_costos(ruta_explicita=None):
    if ruta_explicita:
        if not os.path.exists(ruta_explicita):
            log("ERROR: no existe el Excel de costos: " + ruta_explicita)
            sys.exit(1)
        return ruta_explicita
    candidatos = sorted(glob.glob(os.path.join(BASE_DIR, "PRODUCTOS*.xlsx")),
                        key=os.path.getmtime, reverse=True)
    return candidatos[0] if candidatos else None


def cargar_costos(ruta_excel, campo="costopromedio"):
    try:
        import openpyxl
    except ImportError:
        log("ERROR: falta openpyxl (pip install openpyxl) para leer los costos.")
        sys.exit(1)
    wb = openpyxl.load_workbook(ruta_excel, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = ws.iter_rows(values_only=True)
    enc = list(next(filas))
    idx = {k: enc.index(k) for k in ("productoid", "costo", "costopromedio", "costoultimo")}
    costos = {}
    for f in filas:
        pid = f[idx["productoid"]]
        try:
            pid = int(pid)
        except (TypeError, ValueError):
            continue
        if campo not in idx:
            continue
        val = f[idx[campo]]
        try:
            val = float(val)
        except (TypeError, ValueError):
            val = None
        if val and val > 0:
            costos[pid] = val
    wb.close()
    return costos


# ------------------------------------------------------------------
# Categorias y normalizacion
# ------------------------------------------------------------------


def categoria_web(categoria_posgold):
    cat = sin_acentos((categoria_posgold or "").strip())
    if not cat or cat in CATEGORIAS_NO_VENDIBLES:
        return None
    for origen, slug in CATEGORIAS_POSGOLD.items():
        if sin_acentos(origen) == cat:
            return slug
    for clave, slug in CLAVES_CATEGORIA:
        if clave in cat:
            return slug
    return "jugueteria"


def codigo_limpio(codigo):
    """Los codigos a veces llegan envueltos en % (comodines de POSGold)."""
    c = str(codigo or "").strip().strip("%").strip()
    return c or None


# ------------------------------------------------------------------
# Cache de imagenes (data/mapa_imagenes.json)
# ------------------------------------------------------------------


def cargar_cache():
    if os.path.exists(CACHE_IMAGENES):
        try:
            with open(CACHE_IMAGENES, encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {"actualizado": None, "por_codigo": {}}


def guardar_cache(cache):
    os.makedirs(DIR_DATA, exist_ok=True)
    cache["actualizado"] = dt.datetime.now().isoformat(timespec="seconds")
    with open(CACHE_IMAGENES, "w", encoding="utf-8", newline="\n") as f:
        json.dump(cache, f, ensure_ascii=False, indent=1)


def sembrar_cache_desde_index(cache):
    """Lee el bloque TOP300/CATALOGO_BASE actual del index.html y guarda
    codigo -> {slug, imagen}: es la mejor semilla porque ya paso QA."""
    if not os.path.exists(INDEX):
        return 0
    with open(INDEX, encoding="utf-8") as f:
        html = f.read()
    m = re.search(re.escape(MARCA_INICIO) + r"([\s\S]*?)" + re.escape(MARCA_FIN), html)
    bloque = m.group(1) if m else html[html.find("const CATALOGO_BASE"):html.find("const IMAGENES_ROTAS")]
    if not bloque:
        return 0
    por_codigo = cache.setdefault("por_codigo", {})
    nuevos = 0
    posiciones = [mm.start() for mm in re.finditer(r'id:\s*"([^"]+)"', bloque)]
    for i, pos in enumerate(posiciones):
        fin = posiciones[i + 1] if i + 1 < len(posiciones) else len(bloque)
        trozo = bloque[pos:fin]
        slug = re.match(r'id:\s*"([^"]+)"', trozo).group(1)
        sku = re.search(r'sku:\s*"([^"]*)"', trozo)
        img = re.search(r'image:\s*"([^"]+)"', trozo)
        nom = re.search(r'name:\s*"([^"]*)"', trozo)
        if not (sku and img):
            continue
        codigo = sku.group(1).strip()
        imagen = img.group(1)
        if not codigo or imagen.startswith("data:image/svg"):
            continue
        ent = por_codigo.get(codigo) or {}
        if not ent.get("slug") and slug:
            ent["slug"] = slug
        if not ent.get("imagen") and imagen:
            ent["imagen"] = imagen
        if nom and not ent.get("nombre"):
            ent["nombre"] = nom.group(1)
        if ent:
            por_codigo[codigo] = ent
            nuevos += 1
    return nuevos


def sembrar_cache_desde_dropi(cache):
    """El Excel DROPI enriquecido trae codigo -> URL de foto de la web."""
    candidatos = glob.glob(os.path.join(BASE_DIR, "DROPI_carga_masiva*ENRIQUECIDO*.xlsx"))
    if not candidatos:
        return 0
    try:
        import openpyxl
        wb = openpyxl.load_workbook(candidatos[0], read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        filas = ws.iter_rows(values_only=True)
        enc = list(next(filas))
        i_sku = enc.index("SKU")
        i_img = enc.index("IMAGENES")
        por_codigo = cache.setdefault("por_codigo", {})
        nuevos = 0
        for f in filas:
            sku = str(f[i_sku] or "").strip()
            imgs = str(f[i_img] or "").strip()
            if not sku or not imgs or "storage/products/" not in imgs:
                continue
            url = imgs.split(",")[0].strip()
            if "/full/" in url:
                url = url.replace("/full/", "/thumb/")
            ent = por_codigo.setdefault(sku, {})
            if not ent.get("imagen"):
                ent["imagen"] = url
                nuevos += 1
        wb.close()
        return nuevos
    except Exception as e:
        log("  aviso: no se pudo sembrar desde DROPI: " + str(e)[:120])
        return 0


# ------------------------------------------------------------------
# Resolucion web: slug + foto verificando que el SKU coincida
# ------------------------------------------------------------------


def buscar_slugs_web(env, nombre):
    for consulta in (nombre, " ".join((nombre or "").split()[:3])):
        if not consulta:
            continue
        url = f"{env['WEB_TIENDA']}/catalogo?q={urllib.parse.quote(consulta)}"
        try:
            codigo, html = http_texto(url, timeout=60)
        except Exception:
            continue
        if codigo != 200:
            continue
        slugs = re.findall(r'href="https?://' + re.escape(urllib.parse.urlparse(env["WEB_TIENDA"]).netloc) +
                           r'/producto/([^"#?]+)"', html)
        unicos = list(dict.fromkeys(slugs))[:10]
        if unicos:
            return unicos
        time.sleep(0.12)
    return []


def pagina_producto_web(env, slug):
    url = f"{env['WEB_TIENDA']}/producto/{urllib.parse.quote(slug)}"
    codigo, html = http_texto(url, timeout=60)
    if codigo != 200:
        return None, None, None
    sku = re.search(r"SKU:\s*([A-Za-z0-9%\-\.]+)", html)
    # La foto PROPIA del producto es el primer URL storage del documento
    # (og/JSON-LD/galeria principal aparecen antes de los carruseles de
    # relacionados, que solo traen thumbs de OTROS productos). Antes este
    # regex buscaba el primer /thumb/ y capturaba productos relacionados.
    titulo = re.search(r"<title>([^<]+)</title>", html)
    nombre_pag = titulo.group(1).split(" - ")[0].strip() if titulo else None
    imagen = None
    m = re.search(r'storage/products/(\d+)/(full|thumb)/([^"\'\s)]+)', html)
    if m:
        pid_f, tamano, archivo = m.groups()
        if tamano == "thumb":
            imagen = f"{env['WEB_TIENDA']}/storage/products/{pid_f}/thumb/{archivo}"
        else:
            thumb = f"{env['WEB_TIENDA']}/storage/products/{pid_f}/thumb/{archivo}"
            ok_thumb, _ = imagen_responde(thumb)
            imagen = thumb if ok_thumb else \
                f"{env['WEB_TIENDA']}/storage/products/{pid_f}/full/{archivo}"
    return (sku.group(1).strip() if sku else None), imagen, nombre_pag


def resolver_por_codigo(env, cache, producto, usar_web=True, forzar=False):
    """Devuelve (slug, imagen, origen) para un producto POSGold o None."""
    codigo = producto["codigo"]
    por_codigo = cache.setdefault("por_codigo", {})
    ent = por_codigo.get(codigo) or {}
    if not forzar and ent.get("slug") and ent.get("imagen"):
        return ent["slug"], ent["imagen"], "cache"
    if not usar_web:
        if ent.get("imagen") and ent.get("slug"):
            return ent["slug"], ent["imagen"], "cache"
        return None
    # Imagen ya conocida (DROPI) pero sin slug: basta una busqueda; si hay
    # un unico resultado no hace falta descargar la pagina completa.
    if not forzar and ent.get("imagen") and not ent.get("slug"):
        slugs = buscar_slugs_web(env, producto["nombre"])
        if len(slugs) == 1:
            por_codigo[codigo] = {**ent, "slug": slugs[0]}
            return slugs[0], ent["imagen"], "cache+busqueda"
        if slugs:
            for slug in slugs:
                sku_pag, imagen, _ = pagina_producto_web(env, slug)
                time.sleep(0.08)
                if imagen and sku_pag == codigo:
                    por_codigo[codigo] = {"slug": slug, "imagen": ent["imagen"],
                                          "nombre": producto["nombre"]}
                    return slug, ent["imagen"], "cache+busqueda"
        return None
    slugs = buscar_slugs_web(env, producto["nombre"])
    if not slugs:
        por_codigo[codigo] = {**ent, "slug": ent.get("slug"), "imagen": ent.get("imagen")}
        return None
    # Validar cada candidato: la pagina debe publicar el mismo SKU de POSGold
    # (o, si es unico resultado, al menos el mismo nombre exacto)
    primera = None
    for slug in slugs:
        sku_pag, imagen, nombre_pag = pagina_producto_web(env, slug)
        time.sleep(0.08)
        if not imagen:
            continue
        if primera is None:
            primera = (slug, imagen, nombre_pag)
        if sku_pag == codigo:
            por_codigo[codigo] = {"slug": slug, "imagen": imagen, "nombre": producto["nombre"]}
            return slug, imagen, "web"
    if len(slugs) == 1 and primera:
        slug, imagen, nombre_pag = primera
        if nombre_pag and sin_acentos(nombre_pag) == sin_acentos(producto["nombre"]):
            por_codigo[codigo] = {"slug": slug, "imagen": imagen, "nombre": producto["nombre"]}
            return slug, imagen, "web"
    return None


def resolver_slug_conocido(env, cache, slug, producto):
    """Para destacados: el slug ya lo conocemos, falta la foto."""
    por_codigo = cache.setdefault("por_codigo", {})
    ent = por_codigo.get(producto["codigo"]) or {}
    if ent.get("imagen"):
        return ent["imagen"], "cache"
    sku_pag, imagen, _ = pagina_producto_web(env, slug)
    if imagen:
        por_codigo[producto["codigo"]] = {"slug": slug, "imagen": imagen,
                                           "nombre": producto["nombre"]}
        return imagen, "web"
    return None, None


# ------------------------------------------------------------------
# Generacion del bloque JS
# ------------------------------------------------------------------


def js_texto(valor):
    return json.dumps(valor, ensure_ascii=False)


def generar_bloque_js(env, seleccion, cantidad):
    fecha = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    L = []
    L.append("")
    L.append("        // ==================================================================")
    L.append(f"        // TOP {cantidad} DESDE POSGOLD — generado automaticamente el {fecha}")
    L.append("        // por sync_posgold.py (cron). Son los productos mas rentables que")
    L.append("        // cumplen la regla de la tienda: existencias en bodega + foto real.")
    L.append("        // El producto que deja de cumplir se oculta solo en la proxima")
    L.append("        // corrida. NO editar a mano: este bloque se regenera completo.")
    L.append("        // ==================================================================")
    L.append("        const CATALOGO_BASE = [")
    for p in seleccion:
        precio = p["precio"]
        compare = p["compare"]
        L.append("            {")
        L.append(f"                id: {js_texto(p['slug'])},")
        L.append(f"                name: {js_texto(p['nombre'])},")
        L.append(f"                description: {js_texto('SKU: ' + p['codigo'])},")
        L.append(f"                image: {js_texto(p['imagen'])},")
        L.append(f"                images: [{js_texto(p['imagen'])}],")
        L.append(f"                category: {js_texto(p['categoria_web'])},")
        L.append(f"                sku: {js_texto(p['codigo'])},")
        L.append(f"                url: {js_texto(env['WEB_TIENDA'] + '/producto/' + p['slug'])},")
        L.append(f"                price: {precio},")
        L.append(f"                comparePrice: {compare},")
        L.append("                sizeGroups: [")
        L.append("                    { key: 'unidad', label: 'Unidades', options: [")
        L.append(f"                        {{ id: 'unidad', label: 'Unidad', price: {precio} }}")
        L.append("                    ]}")
        L.append("                ]")
        L.append("            },")
    L.append("        ];")
    # Carrusel "Los imperdibles": los 16 productos con mayor UTILIDAD
    # (plata absoluta que deja cada venta). Viajan SOLO los ids en ese
    # orden — costos y utilidades exactas viven en el reporte local y
    # jamas se publican. El frente arma la fila con los que esten en el
    # catalogo (initTopUtilidadCarousel).
    top_utilidad = sorted(seleccion, key=lambda p: -(p.get("utilidad") or 0))[:16]
    L.append("")
    L.append("        // Carrusel 'Los imperdibles': IDs ordenados por utilidad")
    L.append("        // (plata que deja cada venta). Lo decide el cron con los costos")
    L.append("        // reales; esos costos nunca viajan al sitio publico.")
    L.append("        const TOP_UTILIDAD_IDS = [")
    for p in top_utilidad:
        L.append(f"            {js_texto(p['slug'])},")
    L.append("        ];")
    return "\n".join(L)


def validar_js_nodo(bloque_js):
    """node --check del bloque generado: si hay error de sintaxis, aborta."""
    nodo = shutil.which("node")
    if not nodo:
        log("  aviso: node no disponible, se salta la validacion de sintaxis")
        return True
    ruta = os.path.join(os.environ.get("TEMP", DIR_REPORTES), "jt_top300_check.js")
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    with open(ruta, "w", encoding="utf-8", newline="\n") as f:
        f.write(bloque_js)
    r = subprocess.run([nodo, "--check", ruta], capture_output=True, text=True)
    if r.returncode != 0:
        log("ERROR: sintaxis JS invalida, NO se toca el index:")
        log(r.stderr[:800])
        return False
    return True


def inyectar_en_index(bloque_js, hacer_backup=True):
    with open(INDEX, encoding="utf-8", newline="") as f:
        html = f.read()
    rx = re.compile(re.escape(MARCA_INICIO) + r"[\s\S]*?" + re.escape(MARCA_FIN))
    if not rx.search(html):
        # Fallback: envolver el CATALOGO_BASE legado con los marcadores
        m = re.search(r"^([ \t]*)const CATALOGO_BASE = \[$", html, re.M)
        if not m:
            log("ERROR: el index.html no tiene los marcadores TOP300 ni CATALOGO_BASE.")
            sys.exit(1)
        inicio = m.start()
        m2 = re.search(r"^([ \t]*)\];[ \t]*$", html[m.end():], re.M)
        if not m2:
            log("ERROR: no se encontro el cierre del CATALOGO_BASE.")
            sys.exit(1)
        fin = m.end() + m2.end()
        sangria = m.group(1)
        html = (html[:inicio] + sangria + MARCA_INICIO + "\n" + html[inicio:fin] +
                "\n" + sangria + MARCA_FIN + html[fin:])
    if hacer_backup:
        os.makedirs(DIR_REPORTES, exist_ok=True)
        stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        shutil.copy2(INDEX, os.path.join(DIR_REPORTES, f"index_backup_{stamp}.html"))
        backups = sorted(glob.glob(os.path.join(DIR_REPORTES, "index_backup_*.html")))
        for viejo in backups[:-5]:
            try:
                os.remove(viejo)
            except OSError:
                pass
    reemplazo = MARCA_INICIO + bloque_js + "\n        " + MARCA_FIN
    html_nuevo = rx.sub(lambda _: reemplazo, html, count=1)
    if html_nuevo == html:
        log("ERROR: no se pudo reemplazar el bloque TOP300.")
        sys.exit(1)
    with open(INDEX, "w", encoding="utf-8", newline="") as f:
        f.write(html_nuevo)
    return True


# ------------------------------------------------------------------
# Reporte local (contiene costos: nunca publicar)
# ------------------------------------------------------------------


def escribir_reporte(filas):
    os.makedirs(DIR_REPORTES, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y-%m-%d_%H%M%S")
    ruta = os.path.join(DIR_REPORTES, f"sync_{stamp}.csv")
    campos = ["estado", "motivo", "productoid", "codigo", "nombre", "categoria_posgold",
              "categoria_web", "precio_detal_con_iva", "precio_base_sin_iva", "costo",
              "utilidad", "margen_costo_pct", "margen_precio_pct", "disponible",
              "destacado", "slug_web", "imagen", "origen_imagen"]
    for fila in filas:  # normalizar: la seleccion guarda el slug en "slug"
        if "slug" in fila and not fila.get("slug_web"):
            fila["slug_web"] = fila["slug"]
    with open(ruta, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=campos)
        w.writeheader()
        for fila in filas:
            w.writerow({c: fila.get(c, "") for c in campos})
    return ruta


# ------------------------------------------------------------------
# Publicar en GitHub (solo con --publicar)
# ------------------------------------------------------------------


def publicar_git(env, cantidad):
    if subprocess.run(["git", "rev-parse", "--is-inside-work-tree"],
                      capture_output=True, cwd=BASE_DIR).returncode != 0:
        log("AVISO: esta carpeta no es un repositorio git; no se publico.")
        log("       Conecta el repo (git init + remote add + push) y repite.")
        return False
    stamp = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    for cmd in (
        ["git", "add", "index.html"],
        ["git", "add", "data/mapa_imagenes.json"],
        ["git", "commit", "-m", f"sync: top {cantidad} de POSGold ({stamp})"],
    ):
        r = subprocess.run(cmd, cwd=BASE_DIR, capture_output=True, text=True)
        if r.returncode != 0 and "commit" not in cmd[1]:
            log("AVISO git: " + (r.stderr or r.stdout).strip()[:200])
    # En los runners de GitHub Actions el checkout queda en HEAD separado:
    # se empuja explicitamente HEAD a la rama principal del repo.
    rama = os.environ.get("GITHUB_REF_NAME") or "main"
    r = subprocess.run(["git", "push", "origin", "HEAD:" + rama],
                       cwd=BASE_DIR, capture_output=True, text=True)
    if r.returncode != 0:
        log("AVISO git push: " + (r.stderr or r.stdout).strip()[:300])
        return False
    return True


# ------------------------------------------------------------------
# Proceso principal
# ------------------------------------------------------------------


def main():
    ap = argparse.ArgumentParser(description="Carga el TOP de POSGold al index.html")
    ap.add_argument("--dry-run", action="store_true", help="no escribe el index ni publica")
    ap.add_argument("--publicar", action="store_true", help="git commit + push tras cargar")
    ap.add_argument("--costos", help="ruta del Excel de costos (default: PRODUCTOS*.xlsx mas reciente)")
    ap.add_argument("--limite", type=int, help="cantidad de productos (default: .env TOP300_CANTIDAD)")
    ap.add_argument("--sin-online", action="store_true", help="resolver fotos solo con la cache local")
    ap.add_argument("--sin-verificar-fotos", action="store_true", help="saltar la verificacion HTTP de fotos")
    args = ap.parse_args()

    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    env = cargar_env()
    cantidad = args.limite or int(env.get("TOP300_CANTIDAD", "300"))
    campo_costo = env.get("COSTO_CAMPO", "costopromedio")
    criterio = env.get("RANKING", "margen_costo")
    incluir_destacados = env.get("INCLUIR_DESTACADOS", "true").lower() == "true"
    destacados_ids = [s.strip() for s in env.get("DESTACADOS_IDS", "").split(",") if s.strip()]
    usar_web = not args.sin_online

    log("=" * 66)
    log("SYNC POSGOLD -> index.html | " + dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    log("=" * 66)

    # --- costos ---
    excel = encontrar_excel_costos(args.costos)
    if not excel:
        log("ERROR: no hay Excel de costos (PRODUCTOS *.xlsx) en la carpeta.")
        log("       Exporta el inventario desde POSGold con costo y vuelve a correr.")
        sys.exit(1)
    log(f"\n[1/6] Costos desde: {os.path.basename(excel)} (campo: {campo_costo})")
    costos = cargar_costos(excel, campo_costo)
    log(f"      {len(costos)} productos con costo conocido")

    # --- POSGold ---
    log("\n[2/6] Login POSGold y consulta de productos...")
    jwt = login_jwt(env)
    productos_api = consultar_productos(env, jwt)
    log(f"      {len(productos_api)} productos en bodega {env.get('POSGOLD_BODEGAID')}")

    # --- cache de imagenes ---
    log("\n[3/6] Cache de fotos y slugs...")
    cache = cargar_cache()
    n1 = sembrar_cache_desde_index(cache)
    n2 = sembrar_cache_desde_dropi(cache)
    log(f"      sembrado desde index: {n1} | desde DROPI: {n2} | cache total: {len(cache['por_codigo'])}")

    # --- analisis y filtros ---
    log("\n[4/6] Filtros y rentabilidad...")
    analisis = []
    contadores = {}
    for it in productos_api:
        pid = it.get("Productoid")
        nombre = (it.get("Producto") or "").strip()
        codigo = codigo_limpio(it.get("Codigo") or it.get("Producto_cod"))
        fila = {"productoid": pid, "codigo": codigo or "", "nombre": nombre,
                "categoria_posgold": it.get("Categoria") or "", "destacado": False,
                "disponible": it.get("Disponible") or 0}
        cat = categoria_web(it.get("Categoria"))
        fila["categoria_web"] = cat or ""
        if not cat:
            fila.update({"estado": "OCULTO", "motivo": "categoria no vendible"})
        elif it.get("Activo") is not True:
            fila.update({"estado": "OCULTO", "motivo": "inactivo en POSGold"})
        elif (it.get("Disponible") or 0) <= 0:
            fila.update({"estado": "OCULTO", "motivo": "sin existencias en bodega"})
        elif not it.get("Imagenes") or all(
                PLACEHOLDER_POSGOLD in str(x).lower() for x in it.get("Imagenes")):
            fila.update({"estado": "OCULTO", "motivo": "sin foto en POSGold"})
        elif not (it.get("Precio1") or 0) > 0:
            fila.update({"estado": "OCULTO", "motivo": "sin precio detal"})
        elif pid not in costos:
            fila.update({"estado": "OCULTO", "motivo": f"sin costo en Excel ({campo_costo})"})
        else:
            iva = float(it.get("ProductoImpuestoPorcentaje") or 0)
            p1 = float(it["Precio1"])
            base = p1 / (1 + iva / 100) if iva > 0 else p1
            costo = costos[pid]
            utilidad = base - costo
            fila.update({
                "precio_detal_con_iva": int(round(p1)),
                "precio_base_sin_iva": round(base, 1),
                "costo": costo,
                "utilidad": round(utilidad, 1),
                "margen_costo_pct": round(utilidad / costo * 100, 2) if costo else None,
                "margen_precio_pct": round(utilidad / base * 100, 2) if base else None,
            })
            if utilidad <= 0:
                fila.update({"estado": "OCULTO", "motivo": "rentabilidad negativa"})
            else:
                fila["estado"] = "CANDIDATO"
                fila["motivo"] = ""
        contadores[fila["motivo"] or "candidatos"] = contadores.get(fila["motivo"] or "candidatos", 0) + 1
        analisis.append(fila)

    candidatos = [f for f in analisis if f["estado"] == "CANDIDATO"]
    if criterio == "margen_precio":
        clave = lambda f: f["margen_precio_pct"] or 0
    elif criterio == "utilidad":
        clave = lambda f: f["utilidad"] or 0
    else:
        clave = lambda f: f["margen_costo_pct"] or 0
    candidatos.sort(key=lambda f: (-clave(f), -(f["disponible"] or 0), f["nombre"]))
    log(f"      {len(candidatos)} cumplen todas las reglas (ranking: {criterio})")
    for motivo, n in sorted(contadores.items(), key=lambda kv: -kv[1]):
        log(f"        {n:5}  {motivo}")

    # --- resolucion de fotos y seleccion ---
    log("\n[5/6] Seleccion del TOP " + str(cantidad) + " y resolucion de fotos...")
    por_codigo_api = {}
    for it in productos_api:
        c = codigo_limpio(it.get("Codigo") or it.get("Producto_cod"))
        if c:
            por_codigo_api.setdefault(c, it)

    seleccion = []
    slugs_usados = {}

    # Cortacircuitos: si la web mayorista cae o bloquea al runner, TODAS
    # las verificaciones de foto fallan seguidas. Antes el script se
    # arrastraba por horas (re-resolucion producto por producto con
    # reintentos) hasta que el timeout del workflow lo mataba. Ahora:
    # 25 fallos seguidos = aborto limpio en minutos con diagnostico.
    verificacion = {"fallos_seguidos": 0, "motivos": {}, "muestras": []}
    FALLOS_ABORTAR = 25

    def nota_fallo(motivo, muestra=None):
        verificacion["fallos_seguidos"] += 1
        verificacion["motivos"][motivo] = verificacion["motivos"].get(motivo, 0) + 1
        if muestra and len(verificacion["muestras"]) < 6:
            verificacion["muestras"].append(muestra)

    def abortar_si_web_caida(contexto):
        if verificacion["fallos_seguidos"] < FALLOS_ABORTAR:
            return
        log("")
        log(f"ABORTADO ({contexto}): {verificacion['fallos_seguidos']} intentos "
            f"SEGUIDOS sin poder verificar o encontrar fotos.")
        log("La web mayorista no esta respondiendo desde este entorno")
        log("(sitio caido, en mantenimiento o bloqueando la IP del runner).")
        resumen = ", ".join(f"{m} x{n}" for m, n in
                            sorted(verificacion["motivos"].items(), key=lambda kv: -kv[1]))
        log("Motivos: " + resumen)
        for m in verificacion["muestras"][:6]:
            log("  muestra: " + m)
        log("El index NO se toco: la tienda sigue con la ultima version publicada.")
        log("Reintentar manualmente desde la pestana Actions o esperar la")
        log("siguiente corrida programada.")
        escribir_reporte(analisis)
        sys.exit(1)

    def resolver_y_verificar(fila, forzar=False, slug_fijo=None):
        """Resuelve foto+slug y verifica que la foto responda.
        Solo hace IO: no toca seleccion. Devuelve (slug, imagen, origen)
        o (None, None, motivo)."""
        try:
            # cortacircuitos ya disparado: no iniciar trabajo web nuevo
            if verificacion["fallos_seguidos"] >= FALLOS_ABORTAR and not slug_fijo:
                return None, None, "omitido (web caida)"
            if slug_fijo:
                ent = cache.setdefault("por_codigo", {}).get(fila["codigo"]) or {}
                if ent.get("imagen") and not forzar:
                    imagen, origen = ent["imagen"], "cache"
                else:
                    res_img = resolver_slug_conocido(env, cache, slug_fijo, fila)
                    imagen, origen = (res_img if res_img else (None, None))
                if not imagen:
                    return None, None, "foto sin pagina en la web"
                if not args.sin_verificar_fotos:
                    ok_img, st = imagen_responde(imagen)
                    if not ok_img:
                        return None, None, f"foto rota (HTTP {st})"
                return slug_fijo, imagen, origen
            res = resolver_por_codigo(env, cache, fila, usar_web=usar_web, forzar=forzar)
            if not res:
                return None, None, "foto sin pagina en la web"
            slug, imagen, origen = res
            if not args.sin_verificar_fotos:
                ok_img, st = imagen_responde(imagen)
                if not ok_img:
                    if not forzar and usar_web and verificacion["fallos_seguidos"] < FALLOS_ABORTAR:
                        return resolver_y_verificar(fila, forzar=True)
                    return None, None, f"foto rota (HTTP {st})"
            return slug, imagen, origen
        except Exception as e:
            return None, None, "error de red: " + type(e).__name__

    # destacados primero (si cumplen las mismas reglas)
    if incluir_destacados and destacados_ids:
        slug_a_codigo = {}
        slug_a_nombre = {}
        for k, v in cache.get("por_codigo", {}).items():
            if v.get("slug"):
                slug_a_codigo[v["slug"]] = k
                if v.get("nombre"):
                    slug_a_nombre[v["slug"]] = v["nombre"]
        for slug_dest in destacados_ids:
            fila_dest = None
            codigo_dest = slug_a_codigo.get(slug_dest)
            if codigo_dest and codigo_dest in {f["codigo"] for f in candidatos}:
                fila_dest = next(f for f in candidatos if f["codigo"] == codigo_dest)
            else:
                sku_pag, imagen, nombre_pag = pagina_producto_web(env, slug_dest) if usar_web else (None, None, None)
                if sku_pag and sku_pag in {f["codigo"] for f in candidatos}:
                    fila_dest = next(f for f in candidatos if f["codigo"] == sku_pag)
                elif nombre_pag:
                    fila_dest = next((f for f in candidatos
                                      if sin_acentos(f["nombre"]) == sin_acentos(nombre_pag)), None)
                elif slug_dest in slug_a_nombre:
                    fila_dest = next((f for f in candidatos
                                      if sin_acentos(f["nombre"]) == sin_acentos(slug_a_nombre[slug_dest])), None)
                if fila_dest and imagen:
                    cache.setdefault("por_codigo", {}).setdefault(fila_dest["codigo"], {})["imagen"] = imagen
            if not fila_dest:
                log(f"  destacado fuera: {slug_dest} (no cumple o no esta en POSGold)")
                continue
            slug, imagen, motivo = resolver_y_verificar(fila_dest, slug_fijo=slug_dest)
            if slug:
                verificacion["fallos_seguidos"] = 0
                fila_dest.update({"estado": "ENTRA", "motivo": "", "slug": slug,
                                  "imagen": imagen, "origen_imagen": motivo, "destacado": True})
                if slug not in slugs_usados:
                    slugs_usados[slug] = fila_dest["codigo"]
                    seleccion.append(fila_dest)
                    log(f"  destacado dentro: {fila_dest['nombre'][:40]}")
            else:
                nota_fallo(motivo, f"{motivo} | {fila_dest['nombre'][:36]}")
                abortar_si_web_caida("destacados")
                log(f"  destacado fuera: {slug_dest} ({imagen or motivo})")
        guardar_cache(cache)

    # resto por rentabilidad, en lotes con hilos (el orden del ranking se
    # respeta: los resultados se consumen en el mismo orden en que salieron).
    # El cache se guarda tras cada lote: si la corrida se interrumpe, el
    # trabajo web ya hecho queda y la siguiente corrida es rapida.
    i = 0
    import concurrent.futures as cf
    with cf.ThreadPoolExecutor(max_workers=12) as pool:
        while len(seleccion) < cantidad and i < len(candidatos):
            lote = [f for f in candidatos[i:i + 40] if f.get("estado") == "CANDIDATO"]
            i += 40
            if not lote:
                continue
            # submit + consumo en orden: los fallos se cuentan en vivo y
            # el cortacircuitos puede abortar a mitad del lote
            futuros = [pool.submit(resolver_y_verificar, f) for f in lote]
            for fila, fut in zip(lote, futuros):
                slug, imagen, extra = fut.result()
                if len(seleccion) >= cantidad:
                    break
                if not slug:
                    fila.update({"estado": "OCULTO", "motivo": extra or "foto sin pagina en la web"})
                    nota_fallo(extra or "foto sin pagina en la web",
                               f"{extra} | {fila['nombre'][:36]}")
                    abortar_si_web_caida(f"lote {i // 40}")
                    continue
                verificacion["fallos_seguidos"] = 0
                if slug in slugs_usados:
                    fila.update({"estado": "OCULTO", "motivo": "id duplicado (misma pagina web)"})
                    continue
                fila.update({"estado": "ENTRA", "motivo": "", "slug": slug,
                             "imagen": imagen, "origen_imagen": extra})
                slugs_usados[slug] = fila["codigo"]
                seleccion.append(fila)
            guardar_cache(cache)
            log(f"      lote {i // 40}: {len(seleccion)}/{cantidad} seleccionados "
                f"(candidatos revisados: {min(i, len(candidatos))})")

    # comparePrice: Precio2 sugerido, o precio x 1.19 para el badge
    fallback_xiva = env.get("COMPARE_FALLBACK_XIVA", "true").lower() == "true"
    por_codigo_precio2 = {}
    for it in productos_api:
        c = codigo_limpio(it.get("Codigo") or it.get("Producto_cod"))
        if c:
            por_codigo_precio2[c] = float(it.get("Precio2") or 0)
    for fila in seleccion:
        p1 = fila["precio_detal_con_iva"]
        p2 = por_codigo_precio2.get(fila["codigo"], 0) or 0
        p2 = int(round(p2))
        if p2 > p1:
            fila["compare"] = p2
        else:
            fila["compare"] = int(round(p1 * 1.19)) if fallback_xiva else 0
        fila["precio"] = p1

    # los que cumplen pero no alcanzaron cupo en el TOP quedan marcados
    for fila in analisis:
        if fila.get("estado") == "CANDIDATO":
            fila.update({"estado": "FUERA_DEL_TOP", "motivo": "cumple pero el TOP ya estaba lleno"})

    log(f"\n      seleccion final: {len(seleccion)} productos")
    if len(seleccion) < cantidad:
        log(f"      AVISO: quedaron {cantidad - len(seleccion)} cupos sin llenar "
            f"(productos ocultos por foto/pagina web). Revisa el reporte.")

    # --- generacion y escritura ---
    log("\n[6/6] Generando bloque JS...")
    if not seleccion:
        log("ERROR: ningun producto seleccionado; NO se toca el index.")
        escribir_reporte(analisis)
        sys.exit(1)
    bloque = generar_bloque_js(env, seleccion, len(seleccion))
    if not validar_js_nodo(bloque):
        escribir_reporte(analisis)
        sys.exit(1)

    if args.dry_run:
        log("DRY-RUN: el index NO se modifica. Primer producto del bloque:")
        log("\n".join(bloque.splitlines()[:26]))
    else:
        inyectar_en_index(bloque)
        log("      index.html actualizado (bloque TOP reemplazado)")
        if args.publicar:
            if publicar_git(env, len(seleccion)):
                log("      publicado en GitHub (Pages republica solo)")

    guardar_cache(cache)
    ruta_reporte = escribir_reporte(analisis)
    log(f"\nReporte completo: {ruta_reporte}")

    # resumen de consola
    log("\n" + "=" * 66)
    log("RESUMEN")
    log("=" * 66)
    estados = {}
    for f in analisis:
        k = f["estado"]
        estados[k] = estados.get(k, 0) + 1
    log(f"  productos POSGold analizados : {len(analisis)}")
    for k in sorted(estados, key=lambda x: -estados[x]):
        log(f"  {k:16}: {estados[k]:5}")
    log(f"  top 10 por {criterio}:")
    for f in seleccion[:10]:
        mc = f["margen_costo_pct"] or 0
        log(f"    {mc:6.1f}% margen | ${f['precio_detal_con_iva']:>9,} | {f['nombre'][:42]}")
    sin_foto_web = [f for f in analisis if f.get("motivo") == "foto sin pagina en la web"]
    if sin_foto_web:
        log(f"\n  PENDIENTES: {len(sin_foto_web)} con foto en POSGold pero sin pagina/foto en la web")
        log("  (los primeros 5): " + " | ".join(f["nombre"][:30] for f in sin_foto_web[:5]))
    log("\nListo.")


if __name__ == "__main__":
    main()
